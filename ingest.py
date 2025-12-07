import os
from PyPDF2 import PdfReader
from docx import Document
import openpyxl
from utils.embedding_faiss import EmbeddingManager

class DocumentIngestion:
    def __init__(self):
        """Initialize document ingestion"""
        self.embedding_manager = EmbeddingManager()
        self.supported_formats = ['.pdf', '.docx', '.txt', '.xlsx']
    
    def read_pdf(self, file_path):
        """Extract text from PDF"""
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except Exception as e:
            print(f"Error reading PDF {file_path}: {e}")
            return ""
    
    def read_docx(self, file_path):
        """Extract text from DOCX"""
        try:
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        except Exception as e:
            print(f"Error reading DOCX {file_path}: {e}")
            return ""
    
    def read_txt(self, file_path):
        """Extract text from TXT"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print(f"Error reading TXT {file_path}: {e}")
            return ""
    
    def read_xlsx(self, file_path):
        """Extract text from XLSX"""
        try:
            workbook = openpyxl.load_workbook(file_path)
            text = ""
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) for cell in row if cell is not None]) + "\n"
            return text
        except Exception as e:
            print(f"Error reading XLSX {file_path}: {e}")
            return ""
    
    def chunk_text(self, text, chunk_size=200, overlap=20):
        """Split text into chunks with overlap"""
        words = text.split()
        chunks = []
        
        for i in range(0, len(words), chunk_size - overlap):
            chunk = " ".join(words[i:i + chunk_size])
            if chunk:
                chunks.append(chunk)
        
        return chunks
    
    def ingest_document(self, file_path):
        """Ingest a single document"""
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext not in self.supported_formats:
            print(f"Unsupported format: {file_ext}")
            return []
        
        # Read document based on format
        if file_ext == '.pdf':
            text = self.read_pdf(file_path)
        elif file_ext == '.docx':
            text = self.read_docx(file_path)
        elif file_ext == '.txt':
            text = self.read_txt(file_path)
        elif file_ext == '.xlsx':
            text = self.read_xlsx(file_path)
        else:
            return []
        
        # Chunk the text
        chunks = self.chunk_text(text)
        print(f"Ingested {file_path}: {len(chunks)} chunks")
        
        return chunks
    
    def ingest_directory(self, directory_path):
        """Ingest all documents from a directory"""
        all_chunks = []
        
        for root, dirs, files in os.walk(directory_path):
            for file in files:
                file_path = os.path.join(root, file)
                chunks = self.ingest_document(file_path)
                all_chunks.extend(chunks)
        
        return all_chunks
    
    def build_vector_store(self, directory_path='data/hr_docs'):
        """Build vector store from documents"""
        print(f"Ingesting documents from {directory_path}...")
        
        chunks = self.ingest_directory(directory_path)
        
        if not chunks:
            print("No documents found or ingested")
            return
        
        print(f"Total chunks: {len(chunks)}")
        print("Building FAISS index...")
        
        self.embedding_manager.build_index(chunks)
        self.embedding_manager.save_index()
        
        print("Vector store built successfully!")

if __name__ == "__main__":
    # Example usage
    ingestion = DocumentIngestion()
    ingestion.build_vector_store('data/hr_docs')